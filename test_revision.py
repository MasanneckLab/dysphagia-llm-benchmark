"""Numerical, source-provenance and endpoint tests. Run with the project venv."""
import json
import unittest

import numpy as np
import pandas as pd
from scipy.integrate import quad
from scipy.optimize._numdiff import approx_derivative
from scipy.special import expit
from scipy.stats import norm
from statsmodels.miscmodels.ordinal_model import OrderedModel

from analysis import bootstrap_indices, kappa
from ordinal_mixed import NestedOrdinal, design_matrix
from prepare_data import FILES, HERE, MODELS, binary, normalize, sha256


class RevisionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = pd.read_csv(HERE / 'data/response_level.csv')

    def test_source_files_unchanged(self):
        if not all(p.exists() for p in FILES):
            self.skipTest('Original private working-tree sources are not bundled; verify publication_manifest.json instead.')
        hashes = json.loads((HERE / 'source_audit/source_hashes.json').read_text())
        for path in FILES:
            key = str(path.relative_to(HERE.parents[1]))
            self.assertEqual(hashes[key], sha256(path))

    def test_complete_paired_design(self):
        self.assertEqual(len(self.data), 360)
        self.assertEqual(self.data.item_id.nunique(), 90)
        self.assertEqual(self.data.statement_id.nunique(), 45)
        self.assertTrue(self.data.groupby('statement_id').size().eq(8).all())

    def test_consensus_counts(self):
        expected = dict(zip(MODELS, [66, 59, 55, 51]))
        for m, n in expected.items():
            self.assertEqual(int(self.data.loc[self.data.system == m, 'rating'].eq(2).sum()), n)

    def test_source_support_consensus_denominator(self):
        s = pd.read_csv(HERE / 'results/source_support.csv')
        for row in s.itertuples():
            d = self.data[(self.data.system == row.system) & self.data.rating.eq(2)]
            field = 'guideline_support' if row.endpoint == 'guideline' else 'high_evidence_inclusive'
            self.assertEqual(row.n_correct, len(d))
            self.assertEqual(row.n_supported, int(d[field].eq(1).sum()))
            self.assertAlmostEqual(row.joint_proportion, row.n_supported / 90)
        sonar = s[(s.system == 'Sonar') & (s.endpoint == 'guideline')].iloc[0]
        self.assertEqual((sonar.n_supported, sonar.n_correct), (8, 51))

    def test_audited_workbook_contains_no_formulas(self):
        from openpyxl import load_workbook
        book = load_workbook(HERE / 'data/response_level_audited.xlsx', data_only=False, read_only=True)
        for sheet in book:
            self.assertFalse(any(c.data_type == 'f' for row in sheet.iter_rows() for c in row))
        book.close()

    def test_audited_workbook_roundtrip_preserves_all_fields(self):
        exported = pd.read_excel(HERE / 'data/response_level_audited.xlsx', sheet_name='Response-level data')
        pd.testing.assert_frame_equal(exported, self.data, check_dtype=False, rtol=1e-9, atol=1e-12)

    def test_raw_labels_not_imputed(self):
        d = self.data[(self.data.system == 'Open Evidence') & (self.data.item_id == 76)].iloc[0]
        self.assertEqual(d.high_evidence_BL_raw, 'no')
        self.assertEqual(d.high_evidence_BL, 0)
        self.assertTrue(pd.isna(binary('ni')))
        d = self.data[(self.data.system == 'Open Evidence') & (self.data.item_id == 38)].iloc[0]
        self.assertEqual(d.high_evidence_CH, 0)
        self.assertEqual(d.guideline_support_BL, 1)
        self.assertEqual(d.high_evidence_BL, 0)
        self.assertEqual(d.high_evidence_inclusive, 1)

    def test_master_alignment_and_field_changes(self):
        audit = json.loads((HERE / 'source_audit/master_comparison.json').read_text())
        self.assertEqual(audit['matched_responses'], 360)
        self.assertEqual(audit['master_sha256'], '971d7efb534eb9abbfcdea2e3e05387ccce39f7f1e3286887e939ee7183cff71')
        self.assertTrue(audit['correctness_and_initial_ratings_unchanged'])
        self.assertTrue(audit['CH_source_consensus_unchanged_on_correct_answers'])
        self.assertEqual(len(audit['legacy_formula_cells']), 1)
        self.assertEqual(audit['legacy_formula_cells'][0]['formula'], '=+++M7+E6+E5')
        diff = pd.read_csv(HERE / 'source_audit/master_field_changes.csv')
        self.assertEqual(len(diff), 12)
        self.assertEqual(int(diff.semantic_change.sum()), 10)
        self.assertEqual(set(self.data.source_workbook), {'20260915_Supplementary_Material_3.xlsx'})

    def test_master_reference_repairs_are_explicit(self):
        for system, item, row in [('Sonar', 2, 7), ('Sonar-Reasoning-Pro', 2, 8), ('Sonar', 45, 93)]:
            d = self.data[(self.data.system == system) & (self.data.item_id == item)].iloc[0]
            self.assertEqual(d.source_excel_row, row)
            self.assertNotEqual(d.reference_answer, d.reference_answer_raw)
            self.assertEqual(self.data[self.data.item_id == item].reference_answer.nunique(), 1)
        fragments = self.data[(self.data.item_id == 2) & self.data.system.isin(['Sonar', 'Sonar-Reasoning-Pro'])]
        prefix = fragments[fragments.system == 'Sonar'].iloc[0]
        suffix = fragments[fragments.system == 'Sonar-Reasoning-Pro'].iloc[0]
        self.assertEqual(normalize(prefix.reference_answer_raw) + ' ' + normalize(suffix.reference_answer_raw), normalize(prefix.reference_answer))

    def test_safety_missingness_not_assumed_no(self):
        self.assertEqual(int(self.data.safety_risk_raw.notna().sum()), 1)
        self.assertEqual(int(self.data.safety_risk.eq(0).sum()), 1)
        self.assertEqual(int(self.data.safety_risk.isna().sum()), 359)
        noncorrect = self.data[self.data.rating != 2]
        self.assertEqual((len(noncorrect), int(noncorrect.safety_risk.notna().sum())), (129, 1))

    def test_current_guideline_catalogue(self):
        catalogue = pd.read_csv(HERE / 'data/guideline_titles.csv')
        self.assertEqual(catalogue.guideline_id.tolist(), [1, 2, 3, 4, 5, 6])
        self.assertEqual(sorted(self.data.guideline_id.unique()), [1, 2, 3, 5, 6])

    def test_archived_prompt_matches_collection_hash(self):
        provenance = json.loads((HERE / 'source_audit/api_provenance.json').read_text())
        self.assertEqual(sha256(HERE / 'data/prompt.md'), provenance['gpt_system_prompt_sha256'])

    def test_bootstrap_preserves_all_clusters_and_strata(self):
        draws = bootstrap_indices(100)
        self.assertEqual(draws.shape, (100, 45))
        for i in range(3):
            self.assertTrue(((draws[:, i * 15:(i + 1) * 15] >= 15 * i)
                             & (draws[:, i * 15:(i + 1) * 15] < 15 * (i + 1))).all())

    def test_kappa_reference(self):
        try:
            from sklearn.metrics import cohen_kappa_score
        except ImportError:
            self.skipTest('Optional scikit-learn not installed; formula is tested with hand tables.')
        table = np.zeros((3, 3), dtype=int)
        for r in self.data.itertuples():
            table[r.rater_CH, r.rater_BL] += 1
        for label, weight in [('unweighted', None), ('linear', 'linear'), ('quadratic', 'quadratic')]:
            self.assertAlmostEqual(kappa(table, label), cohen_kappa_score(self.data.rater_CH, self.data.rater_BL, weights=weight), places=12)

    def test_kappa_hand_examples(self):
        self.assertEqual(kappa(np.eye(3)), 1)
        self.assertAlmostEqual(kappa(np.ones((3, 3))), 0)

    def test_analytic_gradient(self):
        X, y, _ = design_matrix(self.data)
        for item_effect in [True, False]:
            model = NestedOrdinal(X, y, 13, item_effect=item_effect)
            tail = [.3, -.25] if item_effect else [.3]
            p = np.r_[-2.3, 1.05, np.arange(X.shape[-1]) * .02, tail]
            numerical = approx_derivative(lambda a: model.objective(a)[0], p, method='3-point')
            np.testing.assert_allclose(model.objective(p)[1], numerical, atol=1e-6, rtol=1e-6)

    def test_no_random_effect_likelihood_matches_statsmodels(self):
        X, y, _ = design_matrix(self.data)
        model = NestedOrdinal(X, y, 15)
        p = np.r_[-2.2, .9, np.arange(X.shape[-1]) * .03, -25, -25]
        fixed = OrderedModel(y.ravel(), X.reshape(360, -1), distr='logit')
        # statsmodels and this implementation use the same transformed threshold gap.
        fixed_ll = fixed.loglike(np.r_[p[2:-2], p[:2]])
        self.assertAlmostEqual(model.objective(p)[0], -fixed_ll, places=9)

    def test_nested_integral_against_independent_adaptive_quadrature(self):
        X = np.tile(np.arange(4)[None, None, :, None] / 3, (1, 2, 1, 1))
        y = np.array([[[0, 1, 2, 2], [1, 1, 2, 0]]])
        p = np.array([-1.8, np.log(2.9), .6, np.log(.9), np.log(.7)])
        model = NestedOrdinal(X, y, 51)
        theta1, theta2 = p[0], p[0] + np.exp(p[1])

        def item_integral(u, item):
            def integrand(v):
                eta = X[0, item, :, 0] * p[2] + u + v
                probs = np.stack([expit(theta1 - eta), expit(theta2 - eta) - expit(theta1 - eta), expit(eta - theta2)], axis=-1)
                return np.prod(probs[np.arange(4), y[0, item]]) * norm.pdf(v, scale=.7)
            return quad(integrand, -7, 7, epsabs=1e-11)[0]

        likelihood = quad(lambda u: item_integral(u, 0) * item_integral(u, 1) * norm.pdf(u, scale=.9),
                          -8, 8, epsabs=1e-11)[0]
        self.assertAlmostEqual(model.objective(p)[0], -np.log(likelihood), places=7)

    def test_numerical_release_gate(self):
        diag = json.loads((HERE / 'results/ordinal_mixed_diagnostics.json').read_text())
        self.assertTrue(diag['release_gate_passed'])
        self.assertEqual(diag['n_responses'], 352)
        self.assertEqual(diag['excluded_separated_guideline_ids'], [5])
        self.assertEqual(diag['sd_item'], 0)
        full = json.loads((HERE / 'results/ordinal_mixed_full_attempt.json').read_text())
        self.assertFalse(full['release_gate_passed'])
        self.assertTrue(full['boundary'])

    def test_marginal_finite_fit_diagnostics(self):
        diagnostics = json.loads((HERE / 'results/ordinal_cluster_robust_diagnostics.json').read_text())
        self.assertEqual(len(diagnostics), 2)
        for d in diagnostics:
            self.assertTrue(d['release_gate_passed'])
            self.assertLess(d['score_max'], 1e-3)
            self.assertGreater(min(d['information_eigenvalues']), 1e-6)
        self.assertEqual(diagnostics[0]['n_responses'], 352)
        self.assertEqual(diagnostics[0]['excluded_separated_guideline_ids'], [5])
        self.assertEqual(diagnostics[1]['n_responses'], 360)
        self.assertFalse(diagnostics[1]['guideline_adjusted'])

    def test_simulated_effect_recovery(self):
        rng = np.random.default_rng(1093)
        # Large calibration dataset: small samples can miss arbitrary error tolerances by chance.
        n = 2000
        X = np.tile(np.array([[0, 0, 0], [1, 0, 0], [0, 1, 0], [0, 0, 1]])[None, None], (n, 2, 1, 1))
        beta = np.array([.8, -.6, .4])
        eta = X @ beta + rng.normal(0, .9, (n, 1, 1)) + rng.normal(0, .6, (n, 2, 1))
        u = rng.uniform(size=(n, 2, 4))
        y = (u > expit(-1.9 - eta)).astype(int) + (u > expit(.7 - eta)).astype(int)
        model = NestedOrdinal(X, y, 19)
        fit = model.fit()
        self.assertTrue(fit.success)
        np.testing.assert_allclose(fit.x[2:5], beta, atol=.23)
        self.assertLess(max(abs(fit.jac)), 1e-3)
        self.assertLess(abs(np.exp(fit.x[-2]) - .9), .3)
        self.assertLess(abs(np.exp(fit.x[-1]) - .6), .3)


if __name__ == '__main__':
    unittest.main(verbosity=2)
